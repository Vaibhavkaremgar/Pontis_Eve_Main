"""Backend tests for the /api/onboarding/parse-resume endpoint (with OCR fallback)."""
import io
import os
from unittest.mock import AsyncMock, patch

import pytest
import requests
from reportlab.pdfgen import canvas
from reportlab.lib.pagesizes import letter
from reportlab.lib.utils import ImageReader
from PIL import Image, ImageDraw, ImageFont
from pypdf import PdfReader

BASE_URL = os.environ.get("REACT_APP_BACKEND_URL").rstrip("/")
PARSE_URL = f"{BASE_URL}/api/onboarding/parse-resume"


@pytest.mark.asyncio
async def test_llm_resume_parse_returns_sanitized_groq_response():
    """A successful Groq response must survive the post-response processing step."""
    import server
    from types import SimpleNamespace

    response = SimpleNamespace(
        choices=[SimpleNamespace(message=SimpleNamespace(content='{"name": "Jane Doe", "skills": ["Python"]}'))]
    )
    with patch.object(
        server.openai_client.chat.completions._pool,
        "chat_completions_create",
        new=AsyncMock(return_value=response),
    ):
        parsed = await server._parse_resume_with_llm("Jane Doe\nPython developer")

    assert parsed["name"] == "Jane Doe"
    assert parsed["skills"] == ["Python"]


@pytest.mark.asyncio
async def test_resume_parse_passes_exact_candidate_correlation_metadata():
    """Replacement parsing is candidate-attributed without parsing PII."""
    import server
    from types import SimpleNamespace

    captured = {}

    async def fake_create(**kwargs):
        captured.update(kwargs)
        return SimpleNamespace(
            choices=[SimpleNamespace(message=SimpleNamespace(content='{"name": "Jane Doe"}'))]
        )

    with patch.object(server.openai_client.chat.completions._pool, "chat_completions_create", new=fake_create):
        await server._parse_resume_with_llm(
            "Jane Doe\nPython developer", candidate_id="candidate-1",
            endpoint="/api/candidate/candidate-1/resume/replace", request_id="opaque-request-1",
        )

    assert captured["_telemetry"] == {
        "workflow": "resume_parsing", "function_name": "_parse_resume_with_llm",
        "candidate_id": "candidate-1", "endpoint": "/api/candidate/candidate-1/resume/replace",
        "request_id": "opaque-request-1",
    }


@pytest.mark.asyncio
async def test_intake_analysis_uses_explicit_chat_endpoint_and_session_id():
    import server
    from types import SimpleNamespace

    captured = {}

    async def fake_create(**kwargs):
        captured.update(kwargs)
        return SimpleNamespace(choices=[SimpleNamespace(message=SimpleNamespace(content="{}"))])

    with patch.object(server.openai_client.chat.completions._pool, "chat_completions_create", new=fake_create):
        await server._llm_analyze_intake(
            {"id": "candidate-1"}, [], None, endpoint="/api/chat", session_id="session-1",
        )

    assert captured["_telemetry"]["endpoint"] == "/api/chat"
    assert captured["_telemetry"]["session_id"] == "session-1"


@pytest.mark.asyncio
async def test_intake_analysis_uses_voice_progress_endpoint_and_call_id():
    import server
    from types import SimpleNamespace

    captured = {}

    async def fake_create(**kwargs):
        captured.update(kwargs)
        return SimpleNamespace(choices=[SimpleNamespace(message=SimpleNamespace(content="{}"))])

    with patch.object(server.openai_client.chat.completions._pool, "chat_completions_create", new=fake_create):
        await server._llm_analyze_intake(
            {"id": "candidate-1"}, [], None,
            endpoint="/api/voice/candidate-intake/progress", vapi_call_id="call-1",
        )

    assert captured["_telemetry"]["endpoint"] == "/api/voice/candidate-intake/progress"
    assert captured["_telemetry"]["vapi_call_id"] == "call-1"


@pytest.mark.asyncio
async def test_resume_persistence_receives_opaque_request_association():
    from unittest.mock import AsyncMock, MagicMock
    import server

    no_duplicate = MagicMock()
    no_duplicate.fetchone.return_value = None
    db = AsyncMock()
    db.execute = AsyncMock(return_value=no_duplicate)
    db.__aenter__ = AsyncMock(return_value=db)
    db.__aexit__ = AsyncMock(return_value=False)
    parsed = {"name": "Candidate", "email": "", "phone": ""}
    upsert = AsyncMock(return_value="candidate-1")
    with patch.object(server, "SessionLocal", return_value=db), \
         patch.object(server, "_extract_resume_text", return_value=("enough text " * 10, False)), \
         patch.object(server, "_parse_resume_with_llm", new=AsyncMock(return_value=parsed)), \
         patch.object(server, "_upsert_candidate", new=upsert), \
         patch.object(server, "_trigger_matching", new=AsyncMock()), \
         patch.object(server, "_seed_employment_gaps_after_parse", new=AsyncMock()):
        await server.parse_resume(file=_make_fake_file(), existing_id=None)

    assert upsert.await_args.kwargs["telemetry_request_id"]
    assert upsert.await_args.kwargs["force_new"] is True


@pytest.mark.asyncio
async def test_voice_extraction_passes_candidate_and_vapi_metadata():
    import server
    from types import SimpleNamespace

    captured = {}

    async def fake_create(**kwargs):
        captured.update(kwargs)
        return SimpleNamespace(choices=[SimpleNamespace(message=SimpleNamespace(content="{}"))])

    with patch.object(server.openai_client.chat.completions._pool, "chat_completions_create", new=fake_create):
        assert await server._extract_voice_info("brief transcript", "candidate-1", "vapi-call-1") == {}

    assert captured["_telemetry"]["candidate_id"] == "candidate-1"
    assert captured["_telemetry"]["vapi_call_id"] == "vapi-call-1"
    assert captured["_telemetry"]["workflow"] == "voice_info_extraction"


@pytest.mark.asyncio
async def test_chat_profile_extraction_passes_candidate_and_session_metadata():
    import server
    from types import SimpleNamespace

    captured = {}

    async def fake_create(**kwargs):
        captured.update(kwargs)
        return SimpleNamespace(choices=[SimpleNamespace(message=SimpleNamespace(content="{}"))])

    with patch.object(server.openai_client.chat.completions._pool, "chat_completions_create", new=fake_create):
        result = await server._extract_multi_field_updates_from_answer(
            "I prefer remote work", ["remote_preference"], "candidate-1", "session-1",
        )

    assert isinstance(result, dict)

    assert captured["_telemetry"]["candidate_id"] == "candidate-1"
    assert captured["_telemetry"]["session_id"] == "session-1"


def _build_text_pdf(lines):
    buf = io.BytesIO()
    c = canvas.Canvas(buf)
    y = 800
    for line in lines:
        c.drawString(50, y, line)
        y -= 18
    c.showPage()
    c.save()
    buf.seek(0)
    return buf.getvalue()


def _build_empty_pdf():
    buf = io.BytesIO()
    c = canvas.Canvas(buf)
    c.showPage()
    c.save()
    buf.seek(0)
    return buf.getvalue()


def _build_image_pdf(text):
    """Render text into a PNG, then draw the PNG on a PDF page.
    Result: a PDF whose page content stream has NO text objects — only an image.
    pypdf.extract_text should therefore return ~empty; OCR fallback should kick in.
    """
    img = Image.new("RGB", (1200, 300), "white")
    draw = ImageDraw.Draw(img)
    try:
        font = ImageFont.truetype(
            "/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf", 40
        )
    except Exception:
        font = ImageFont.load_default()
    draw.text((40, 40), text, fill="black", font=font)
    # Add a second line to ensure enough OCR characters
    draw.text(
        (40, 120),
        "Senior Product Designer at Acme Corp since 2020.",
        fill="black",
        font=font,
    )
    draw.text(
        (40, 200),
        "Skills: Figma, Design Systems, UX Research, Accessibility.",
        fill="black",
        font=font,
    )

    img_buf = io.BytesIO()
    img.save(img_buf, format="PNG")
    img_buf.seek(0)

    buf = io.BytesIO()
    c = canvas.Canvas(buf, pagesize=letter)
    c.drawImage(ImageReader(img_buf), 40, 400, width=520, height=130)
    c.showPage()
    c.save()
    buf.seek(0)
    return buf.getvalue()


# --- Rejects unsupported files ---
def test_reject_unsupported_extension():
    files = {"file": ("resume.txt", b"just some text", "text/plain")}
    r = requests.post(PARSE_URL, files=files, timeout=30)
    assert r.status_code == 400
    detail = r.json().get("detail", "")
    assert "PDF" in detail and "DOC" in detail and "DOCX" in detail


def test_extract_docx_text():
    from docx import Document
    import server

    document = Document()
    document.add_paragraph("Jane Doe")
    document.add_paragraph("Senior Product Designer with Python experience.")
    output = io.BytesIO()
    document.save(output)

    text, used_ocr = server._extract_resume_text(output.getvalue(), ".docx")
    assert "Jane Doe" in text
    assert "Python" in text
    assert used_ocr is False


def test_extract_doc_text_uses_legacy_word_reader():
    import server
    from types import SimpleNamespace
    from unittest.mock import patch

    completed = SimpleNamespace(returncode=0, stdout=b"Jane Doe\nLegacy Word resume", stderr=b"")
    with patch.object(server.subprocess, "run", return_value=completed) as run:
        text, used_ocr = server._extract_resume_text(b"legacy doc", ".doc")

    run.assert_called_once()
    assert text == "Jane Doe\nLegacy Word resume"
    assert used_ocr is False


# --- Rejects empty PDFs with updated error message ---
def test_empty_pdf_rejected_with_updated_message():
    pdf_bytes = _build_empty_pdf()
    files = {"file": ("empty.pdf", pdf_bytes, "application/pdf")}
    r = requests.post(PARSE_URL, files=files, timeout=60)
    assert r.status_code == 400
    detail = r.json().get("detail", "")
    assert "Resume appears empty" in detail
    assert "text-based PDF or a clearer scan" in detail


# --- Happy path: real text-based PDF, used_ocr must be False ---
def test_parse_text_pdf_used_ocr_false():
    lines = [
        "Jane Doe",
        "Email: jane.doe@example.com",
        "Location: San Francisco, USA",
        "",
        "Summary:",
        "Senior Product Designer with 8 years of experience in B2B SaaS.",
        "",
        "Experience:",
        "Senior Product Designer, Acme Corp (2020 - Present)",
        "Led design system across 4 product teams.",
        "Product Designer, Beta Inc (2016 - 2020)",
        "",
        "Education:",
        "B.A. in Design, Stanford University (2012 - 2016)",
        "",
        "Skills: Figma, Design Systems, UX Research, Prototyping, Accessibility",
    ]
    pdf_bytes = _build_text_pdf(lines)
    files = {"file": ("jane.pdf", pdf_bytes, "application/pdf")}
    r = requests.post(PARSE_URL, files=files, timeout=90)
    assert r.status_code == 200, r.text
    data = r.json()
    for key in ["name", "email", "headline", "location", "bio"]:
        assert key in data and isinstance(data[key], str)
    for key in ["experience", "education", "keySkills"]:
        assert key in data and isinstance(data[key], list)
    assert "_meta" in data
    assert data["_meta"].get("used_ocr") is False


# --- OCR fallback: image-based PDF ---
def test_image_pdf_triggers_ocr_fallback():
    pdf_bytes = _build_image_pdf("Alice Smith - Product Designer")

    # Sanity check locally: pypdf should extract <40 chars from this PDF
    reader = PdfReader(io.BytesIO(pdf_bytes))
    local_text = "\n".join((p.extract_text() or "") for p in reader.pages).strip()
    assert len(local_text) < 40, f"Expected image-only PDF, got extracted text: {local_text!r}"

    files = {"file": ("scanned.pdf", pdf_bytes, "application/pdf")}
    r = requests.post(PARSE_URL, files=files, timeout=180)
    assert r.status_code == 200, r.text
    data = r.json()
    assert "_meta" in data
    assert data["_meta"].get("used_ocr") is True, f"Expected OCR fallback, meta={data.get('_meta')}"


# ---------------------------------------------------------------------------
# Regression tests: mismatch-first, duplicate-after-match, normal upload
# ---------------------------------------------------------------------------

def _make_fake_db(fetchone_return=None):
    """Build a minimal async context-manager DB mock."""
    from unittest.mock import AsyncMock, MagicMock
    fake_result = MagicMock()
    fake_result.fetchone.return_value = fetchone_return
    fake_db = AsyncMock()
    fake_db.execute = AsyncMock(return_value=fake_result)
    fake_db.__aenter__ = AsyncMock(return_value=fake_db)
    fake_db.__aexit__ = AsyncMock(return_value=False)
    return fake_db


def _make_fake_file(filename="resume.pdf", content=b"%PDF-1.4 fake"):
    from unittest.mock import AsyncMock, MagicMock
    from fastapi import UploadFile
    f = MagicMock(spec=UploadFile)
    f.filename = filename
    f.read = AsyncMock(return_value=content)
    return f


@pytest.mark.asyncio
async def test_identity_mismatch_email_stops_before_duplicate_check():
    """
    When the resume email differs from the logged-in candidate's email,
    a 422 must be raised immediately — duplicate-resume check must NOT run.
    """
    from unittest.mock import AsyncMock, MagicMock, patch, call
    import server
    from fastapi import HTTPException

    # DB row for the logged-in candidate (email mismatch)
    id_row = MagicMock()
    id_row.__getitem__ = lambda self, i: ("real@example.com" if i == 0 else "")
    id_result = MagicMock()
    id_result.fetchone.return_value = id_row

    fake_db = AsyncMock()
    fake_db.execute = AsyncMock(return_value=id_result)
    fake_db.__aenter__ = AsyncMock(return_value=fake_db)
    fake_db.__aexit__ = AsyncMock(return_value=False)

    with patch.object(server, "SessionLocal", return_value=fake_db), \
         patch.object(server, "_extract_pdf_text", return_value=("enough text " * 10, False)), \
         patch.object(server, "_parse_resume_with_llm",
                      new=AsyncMock(return_value={"name": "Other", "email": "other@example.com", "phone": ""})):

        with pytest.raises(HTTPException) as exc_info:
            await server.parse_resume(file=_make_fake_file(), existing_id="some-candidate-uuid")

    assert exc_info.value.status_code == 422
    assert "email" in exc_info.value.detail.lower()
    # Duplicate check (fingerprint query) must NOT have been called
    for c in fake_db.execute.call_args_list:
        sql = str(c)
        assert "resume_fingerprint" not in sql, "Duplicate check ran despite identity mismatch"


@pytest.mark.asyncio
async def test_identity_mismatch_phone_stops_before_duplicate_check():
    """
    When the resume phone differs from the logged-in candidate's phone,
    a 422 must be raised — duplicate check must NOT run.
    """
    from unittest.mock import AsyncMock, MagicMock, patch
    import server
    from fastapi import HTTPException

    id_row = MagicMock()
    id_row.__getitem__ = lambda self, i: ("" if i == 0 else "+1-555-000-0001")
    id_result = MagicMock()
    id_result.fetchone.return_value = id_row

    fake_db = AsyncMock()
    fake_db.execute = AsyncMock(return_value=id_result)
    fake_db.__aenter__ = AsyncMock(return_value=fake_db)
    fake_db.__aexit__ = AsyncMock(return_value=False)

    with patch.object(server, "SessionLocal", return_value=fake_db), \
         patch.object(server, "_extract_pdf_text", return_value=("enough text " * 10, False)), \
         patch.object(server, "_parse_resume_with_llm",
                      new=AsyncMock(return_value={"name": "Other", "email": "", "phone": "+1-555-999-9999"})):

        with pytest.raises(HTTPException) as exc_info:
            await server.parse_resume(file=_make_fake_file(), existing_id="some-candidate-uuid")

    assert exc_info.value.status_code == 422
    assert "mobile" in exc_info.value.detail.lower() or "phone" in exc_info.value.detail.lower()
    for c in fake_db.execute.call_args_list:
        assert "resume_fingerprint" not in str(c), "Duplicate check ran despite phone mismatch"


@pytest.mark.asyncio
async def test_duplicate_resume_returns_409_after_identity_match():
    """
    When identity matches (or no existing_id), a duplicate fingerprint must
    return HTTP 409 with the correct message.
    """
    from unittest.mock import AsyncMock, MagicMock, patch
    import server
    from fastapi import HTTPException

    fake_fingerprint = "aabbcc" * 10

    # First DB call (identity check for existing_id): returns matching candidate
    id_row = MagicMock()
    id_row.__getitem__ = lambda self, i: ("same@example.com" if i == 0 else "")
    id_result = MagicMock()
    id_result.fetchone.return_value = id_row

    # Second DB call (fingerprint check): returns a row (duplicate found)
    dup_row = MagicMock()
    dup_result = MagicMock()
    dup_result.fetchone.return_value = dup_row

    call_count = 0

    async def _execute_side_effect(*args, **kwargs):
        nonlocal call_count
        call_count += 1
        return id_result if call_count == 1 else dup_result

    fake_db = AsyncMock()
    fake_db.execute = AsyncMock(side_effect=_execute_side_effect)
    fake_db.__aenter__ = AsyncMock(return_value=fake_db)
    fake_db.__aexit__ = AsyncMock(return_value=False)

    with patch.object(server, "SessionLocal", return_value=fake_db), \
         patch("hashlib.sha256") as mock_sha, \
         patch.object(server, "_extract_pdf_text", return_value=("enough text " * 10, False)), \
         patch.object(server, "_parse_resume_with_llm",
                      new=AsyncMock(return_value={"name": "Same", "email": "same@example.com", "phone": ""})):
        mock_sha.return_value.hexdigest.return_value = fake_fingerprint

        with pytest.raises(HTTPException) as exc_info:
            await server.parse_resume(file=_make_fake_file(), existing_id="some-candidate-uuid")

    assert exc_info.value.status_code == 409
    assert "Duplicate resume" in exc_info.value.detail


@pytest.mark.asyncio
async def test_normal_upload_no_existing_id_duplicate_check_runs():
    """
    When no existing_id is provided (new candidate), identity check is skipped
    and duplicate fingerprint check runs. A duplicate must return 409.
    """
    from unittest.mock import AsyncMock, MagicMock, patch
    import server
    from fastapi import HTTPException

    fake_fingerprint = "ccddee" * 10

    dup_row = MagicMock()
    dup_result = MagicMock()
    dup_result.fetchone.return_value = dup_row

    fake_db = AsyncMock()
    fake_db.execute = AsyncMock(return_value=dup_result)
    fake_db.__aenter__ = AsyncMock(return_value=fake_db)
    fake_db.__aexit__ = AsyncMock(return_value=False)

    with patch.object(server, "SessionLocal", return_value=fake_db), \
         patch("hashlib.sha256") as mock_sha, \
         patch.object(server, "_extract_pdf_text", return_value=("enough text " * 10, False)), \
         patch.object(server, "_parse_resume_with_llm",
                      new=AsyncMock(return_value={"name": "New", "email": "new@example.com", "phone": ""})):
        mock_sha.return_value.hexdigest.return_value = fake_fingerprint

        with pytest.raises(HTTPException) as exc_info:
            await server.parse_resume(file=_make_fake_file(), existing_id=None)

    assert exc_info.value.status_code == 409
    assert "Duplicate resume" in exc_info.value.detail


@pytest.mark.asyncio
async def test_failed_resume_parsing_does_not_reach_candidate_association():
    from unittest.mock import AsyncMock, patch
    import server

    with patch.object(server, "_extract_resume_text", return_value=("enough text " * 10, False)), \
         patch.object(server, "_parse_resume_with_llm", new=AsyncMock(side_effect=RuntimeError("parse failed"))), \
         patch.object(server, "_upsert_candidate", new=AsyncMock()) as upsert:
        with pytest.raises(RuntimeError, match="parse failed"):
            await server.parse_resume(file=_make_fake_file(), existing_id=None)

    upsert.assert_not_awaited()


# --- Unit test: Groq 429 rate-limit → endpoint returns 429, not 500 ---
@pytest.mark.asyncio
async def test_parse_resume_llm_rate_limit_returns_429():
    """_parse_resume_with_llm must raise HTTPException(429) on RateLimitError,
    so the endpoint returns 429 instead of an unhandled 500."""
    from openai import RateLimitError
    from fastapi import HTTPException
    import server

    fake_response = type("R", (), {"status_code": 429, "headers": {}, "text": "rate limited"})()
    rate_limit_exc = RateLimitError("rate limited", response=fake_response, body=None)

    with patch.object(server.openai_client.chat.completions, "create", new=AsyncMock(side_effect=rate_limit_exc)):
        with pytest.raises(HTTPException) as exc_info:
            await server._parse_resume_with_llm("some resume text")

    assert exc_info.value.status_code == 429
    assert "temporarily unavailable" in exc_info.value.detail.lower() or "rate" in exc_info.value.detail.lower()
